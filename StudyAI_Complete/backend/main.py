import os
import base64
import io
import json
import re
import time
import sqlite3
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.responses import HTMLResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from google import genai
from google.genai import types

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None
try:
    from docx import Document
except Exception:
    Document = None
try:
    from pptx import Presentation
except Exception:
    Presentation = None

app = FastAPI(title="Study AI Backend", version="4.0")

origins = [x.strip() for x in os.getenv("ALLOWED_ORIGINS", "*").split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins or ["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

API_KEY = os.getenv("GEMINI_API_KEY")
MODEL = (os.getenv("GEMINI_MODEL") or "gemini-2.5-flash-lite").strip()
client = genai.Client(api_key=API_KEY) if API_KEY else None

# Large-file protection: keep each AI request comfortably below the TPM limit.
CHUNK_CHARS = int(os.getenv("AI_CHUNK_CHARS", "8000"))
LARGE_FILE_CHARS = int(os.getenv("AI_LARGE_FILE_CHARS", "16000"))
MAX_RETRIES = 0
AI_TIMEOUT_SECONDS = 45


# Public app mode: no device activation or license gate.
MAX_UPLOAD_BYTES = int(os.getenv('MAX_UPLOAD_BYTES', str(15 * 1024 * 1024)))
RATE_WINDOW_SECONDS = int(os.getenv('RATE_WINDOW_SECONDS', '60'))
RATE_MAX_REQUESTS = int(os.getenv('RATE_MAX_REQUESTS', '30'))
_rate_hits = {}

def _client_ip(request: Request) -> str:
    return (request.client.host if request.client else 'unknown')

def rate_limit(request: Request):
    now = time.time(); ip = _client_ip(request)
    hits = [t for t in _rate_hits.get(ip, []) if now - t < RATE_WINDOW_SECONDS]
    if len(hits) >= RATE_MAX_REQUESTS:
        raise HTTPException(429, 'تم تجاوز عدد الطلبات المؤقت المسموح به. حاول بعد قليل.')
    hits.append(now); _rate_hits[ip] = hits
    if len(_rate_hits) > 2000:
        for k, v in list(_rate_hits.items()):
            if not v or now - v[-1] > RATE_WINDOW_SECONDS:
                _rate_hits.pop(k, None)

def ensure_upload_size(data: bytes):
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f'حجم الملف كبير جداً. الحد الأقصى {MAX_UPLOAD_BYTES // (1024*1024)} MB.')


# Prototype storage for one user's latest uploaded material.
# For a multi-user production service, replace this with per-user/session storage.
last_material = {
    "name": None,
    "text": "",
    "image_data": None,
    "mime": None,
}

SYSTEM = """You are Study AI, an academic study assistant specialized in medical and pharmaceutical education.

Core rules:
- If a study material is explicitly provided for the current request, use it as the primary source.
- If no material is provided, answer the user normally using your general knowledge, while keeping medical/pharmaceutical answers educational and careful.
- Do not invent scientific facts, doses, drug names, numbers, units, abbreviations, or exam certainty.
- If the material is unclear or missing the answer, say so.
- If you add knowledge not present in the material, label it exactly as: "معلومة إضافية".
- For Arabic translation, use accurate medical/pharmaceutical terminology and keep the English term beside it when useful.
- Keep the material's meaning; do not silently omit important information.
- Prefer clear Arabic RTL-friendly headings and concise bullets.
- Never claim that a question is guaranteed to appear on an exam. Use "High-Yield" or "مرجّح للمراجعة".
"""

def require_client():
    if not client:
        raise HTTPException(503, "GEMINI_API_KEY غير مضبوط على الخادم.")

def extract_text(data: bytes, name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf" and PdfReader:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx" and Document:
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == ".pptx" and Presentation:
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text)
        return "\n".join(parts)
    if ext in [".txt", ".md"]:
        return data.decode("utf-8", errors="replace")
    return ""

def clean_json(raw: str):
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.I)
        raw = re.sub(r"\s*```$", "", raw)
    start = raw.find("{")
    end = raw.rfind("}")
    if start >= 0 and end > start:
        raw = raw[start:end + 1]
    return json.loads(raw)

def make_input(instruction: str):
    if last_material["image_data"]:
        return [{
            "role": "user",
            "content": [
                {"type": "input_text", "text": instruction},
                {
                    "type": "input_image",
                    "image_url": f"data:{last_material['mime']};base64,{last_material['image_data']}",
                },
            ],
        }]
    return instruction + "\n\nSTUDY MATERIAL:\n" + last_material["text"]

def _gemini_generate(contents, max_output_tokens=2500, schema=None):
    require_client()
    config_kwargs = {
        "system_instruction": SYSTEM,
        "max_output_tokens": max_output_tokens,
        "temperature": 0.2,
    }
    if schema is not None:
        config_kwargs["response_mime_type"] = "application/json"
        config_kwargs["response_schema"] = schema
    try:
        response = client.models.generate_content(
            model=MODEL,
            contents=contents,
            config=types.GenerateContentConfig(**config_kwargs),
        )
        text = (getattr(response, "text", "") or "").strip()
        if not text:
            raise RuntimeError("GEMINI_EMPTY_RESPONSE")
        return text
    except Exception as exc:
        msg = str(exc)
        low = msg.lower()
        if "401" in low or "403" in low or "api key" in low or "authentication" in low or "permission" in low:
            raise RuntimeError("GEMINI_AUTH_ERROR") from exc
        if "429" in low or "quota" in low or "rate limit" in low or "resource exhausted" in low:
            raise RuntimeError("GEMINI_RATE_LIMIT") from exc
        if "timeout" in low or "timed out" in low:
            raise RuntimeError("GEMINI_TIMEOUT") from exc
        raise RuntimeError("GEMINI_REQUEST_ERROR") from exc

def ask_ai(instruction: str, text_override: str | None = None, schema=None, name="study_analysis", max_output_tokens=2500) -> str:
    require_client()
    if text_override is None:
        input_value = make_input(instruction)
    else:
        input_value = instruction + "\n\nSTUDY MATERIAL EXCERPT:\n" + text_override
    return _gemini_generate(input_value, max_output_tokens=max_output_tokens, schema=schema)


@app.get("/health")
def health():
    return {
        "ok": True,
        "model": MODEL,
        "ai_configured": bool(API_KEY),
        "material_loaded": bool(last_material["name"]),
        "activation_required": False,
    }



class AITestRequest(BaseModel):
    question: str = "Reply with exactly: OK"

@app.post("/ai/test")
def ai_test(body: AITestRequest, request: Request):
    rate_limit(request)
    q=(body.question or "Reply with exactly: OK").strip()[:200]
    if not API_KEY:
        raise HTTPException(503, "مفتاح Gemini غير مضبوط على Render.")
    try:
        answer=ask_ai_with_optional_material(
            "Perform a connectivity test. Reply briefly and do not use any study material. User request:\n"+q,
            max_output_tokens=24
        )
        return {"ok":True,"answer":answer,"model":MODEL}
    except Exception as exc:
        code=str(exc)
        if code == "GEMINI_AUTH_ERROR":
            raise HTTPException(502,"مفتاح Gemini غير صالح أو منتهي على Render.")
        if code == "GEMINI_RATE_LIMIT":
            raise HTTPException(429,"وصلنا إلى حد Gemini المجاني/المؤقت. انتظر قليلاً ثم جرّب مرة واحدة.")
        if code == "GEMINI_TIMEOUT":
            raise HTTPException(504,"Gemini لم يرد خلال المهلة. جرّب مرة واحدة بعد قليل.")
        raise HTTPException(502,"تعذر الوصول إلى Gemini حالياً. راجع سجل Render لمعرفة السبب.")

@app.post("/analyze")
async def analyze(request: Request, file: UploadFile = File(...)):
    rate_limit(request)
    data = await file.read()
    ensure_upload_size(data)
    mime = file.content_type or ""
    text = extract_text(data, file.filename)

    if not text and not mime.startswith("image/"):
        raise HTTPException(
            400,
            "لم أستطع استخراج النص من هذا الملف. استخدم PDF/DOCX/PPTX أو صورة واضحة."
        )

    last_material["name"] = file.filename
    last_material["text"] = text
    last_material["mime"] = mime
    last_material["image_data"] = (
        base64.b64encode(data).decode("ascii") if mime.startswith("image/") else None
    )

    prompt = """حلّل المادة الدراسية تحليلاً شاملاً، ثم أعد النتيجة كـ JSON صالح فقط، بدون Markdown وبدون أي نص خارج JSON.
استخدم هذه المفاتيح حرفياً:
{
  "title": "اسم المادة أو الموضوع",
  "summary": "ملخص واضح",
  "difficulty": "سهل/متوسط/صعب",
  "understanding": 0,
  "high_yield": ["نقاط High-Yield"],
  "topics": ["المواضيع الرئيسية"],
  "must_memorize": ["ما يجب حفظه"],
  "must_understand": ["ما يجب فهمه"],
  "terms": [{"english":"", "arabic":"", "explanation":""}],
  "comparisons": ["مقارنات مهمة"],
  "rewrite": "A clear, organized English rewrite of the study material, preserving the original scientific meaning and important terminology.",
  "study_plan": ["خطوات مرتبة لدراسة المادة"],
  "quick_revision": ["نقاط مراجعة سريعة"],
  "questions": [{"question":"", "options":["A) ","B) ","C) ","D) "], "answer":"", "explanation":""}],
  "flashcards": [{"front":"", "back":""}]
}
اجعل understanding رقماً من 0 إلى 100 يمثل وضوح المادة واستيعاب بنيتها، وليس تشخيصاً لمستوى الطالب.
لا تضع معلومات غير موجودة في المادة إلا إذا وسمتها داخل النص بأنها "معلومة إضافية".
"""

    full_schema = {
        "type": "object", "additionalProperties": False,
        "properties": {
            "title": {"type": "string"}, "summary": {"type": "string"}, "difficulty": {"type": "string"},
            "understanding": {"type": "integer", "minimum": 0, "maximum": 100},
            "high_yield": {"type": "array", "items": {"type": "string"}},
            "topics": {"type": "array", "items": {"type": "string"}},
            "must_memorize": {"type": "array", "items": {"type": "string"}},
            "must_understand": {"type": "array", "items": {"type": "string"}},
            "terms": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {"english": {"type": "string"}, "arabic": {"type": "string"}, "explanation": {"type": "string"}}, "required": ["english", "arabic", "explanation"]}},
            "comparisons": {"type": "array", "items": {"type": "string"}}, "rewrite": {"type": "string"},
            "study_plan": {"type": "array", "items": {"type": "string"}}, "quick_revision": {"type": "array", "items": {"type": "string"}},
            "questions": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {"question": {"type": "string"}, "options": {"type": "array", "items": {"type": "string"}}, "answer": {"type": "string"}, "explanation": {"type": "string"}}, "required": ["question", "options", "answer", "explanation"]}},
            "flashcards": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {"front": {"type": "string"}, "back": {"type": "string"}}, "required": ["front", "back"]}}
        },
        "required": ["title", "summary", "difficulty", "understanding", "high_yield", "topics", "must_memorize", "must_understand", "terms", "comparisons", "rewrite", "study_plan", "quick_revision", "questions", "flashcards"]
    }

    chunk_schema = {
        "type": "object", "additionalProperties": False,
        "properties": {
            "title": {"type": "string"}, "summary": {"type": "string"},
            "high_yield": {"type": "array", "items": {"type": "string"}},
            "topics": {"type": "array", "items": {"type": "string"}},
            "must_memorize": {"type": "array", "items": {"type": "string"}},
            "must_understand": {"type": "array", "items": {"type": "string"}},
            "terms": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {"english": {"type": "string"}, "arabic": {"type": "string"}, "explanation": {"type": "string"}}, "required": ["english", "arabic", "explanation"]}},
            "comparisons": {"type": "array", "items": {"type": "string"}}
        },
        "required": ["title", "summary", "high_yield", "topics", "must_memorize", "must_understand", "terms", "comparisons"]
    }

    # Small/medium files: one request, with a hard input cap that stays below the rate limit.
    if len(text) <= LARGE_FILE_CHARS:
        try:
            raw = ask_ai(prompt, text_override=text[:LARGE_FILE_CHARS], schema=full_schema, name="study_analysis", max_output_tokens=5000)
            result = json.loads(raw)
            return {"filename": file.filename, "model": MODEL, "analysis": result}
        except Exception as exc:
            if "rate_limit" in str(exc).lower() or "429" in str(exc):
                raise HTTPException(429, "تم تجاوز حد Gemini مؤقتاً. انتظر قليلاً ثم أعد المحاولة.")
            raise HTTPException(502, f"فشل تحليل المادة: {exc}")

    # Large files: analyze sequential chunks, then synthesize the compact chunk results.
    chunks = []
    pos = 0
    while pos < len(text):
        end = min(len(text), pos + CHUNK_CHARS)
        chunks.append(text[pos:end])
        if end >= len(text):
            break
        pos = max(end - 1200, pos + 1)  # small overlap keeps definitions across boundaries

    partials = []
    for i, chunk in enumerate(chunks, 1):
        instruction = f"حلّل الجزء {i} من {len(chunks)} من المادة فقط. استخرج المعلومات المهمة الموجودة فعلياً في هذا الجزء، ولا تخترع معلومات. كن مختصراً لأن هذه نتيجة وسيطة ستُدمج لاحقاً."
        try:
            raw = ask_ai(instruction, text_override=chunk, schema=chunk_schema, name="study_chunk", max_output_tokens=2200)
            partials.append(json.loads(raw))
        except Exception as exc:
            if "rate_limit" in str(exc).lower() or "429" in str(exc):
                raise HTTPException(429, "الملف كبير جداً على حد Gemini الحالي. أعد المحاولة بعد دقيقة.")
            raise HTTPException(502, f"فشل تحليل الجزء {i}: {exc}")
        time.sleep(1.2)

    combined = json.dumps(partials, ensure_ascii=False)
    final_instruction = """ادمج نتائج تحليل أجزاء المادة التالية في تحليل نهائي واحد. لا تضف معلومات غير موجودة في نتائج الأجزاء. احذف التكرار، وحافظ على أهم النقاط والمصطلحات، ثم أنشئ النتيجة الكاملة بالمخطط المطلوب.\n\nنتائج الأجزاء:\n""" + combined
    try:
        raw = ask_ai(final_instruction, text_override="", schema=full_schema, name="study_analysis", max_output_tokens=2600)
        result = json.loads(raw)
    except Exception as exc:
        if "rate_limit" in str(exc).lower() or "429" in str(exc):
            raise HTTPException(429, "تم تجاوز حد Gemini مؤقتاً أثناء تجميع التحليل. أعد المحاولة بعد دقيقة.")
        raise HTTPException(502, f"فشل تجميع التحليل: {exc}")

    return {"filename": file.filename, "model": MODEL, "analysis": result}



def ask_feature(instruction: str, max_output_tokens: int = 1600) -> str:
    """Run feature requests with bounded context to reduce TPM bursts."""
    text = last_material.get("text", "")
    if not text:
        return ask_ai(instruction, max_output_tokens=max_output_tokens)

    # Rewrite/translation can be assembled from smaller sections.
    if len(text) <= 18000:
        return ask_ai(instruction, text_override=text, max_output_tokens=max_output_tokens)

    parts = []
    step = 12000
    overlap = 600
    pos = 0
    while pos < len(text):
        end = min(len(text), pos + step)
        chunk = text[pos:end]
        parts.append(ask_ai(
            instruction + f"\n\nهذا الجزء {len(parts)+1} من المادة. عالج هذا الجزء فقط وحافظ على المصطلحات المهمة.",
            text_override=chunk,
            max_output_tokens=max_output_tokens,
        ))
        if end >= len(text):
            break
        pos = max(end - overlap, pos + 1)
        time.sleep(0.8)

    return "\n\n---\n\n".join(parts)

FEATURE_PROMPTS = {
    "rewrite": """Rewrite the uploaded study material in clear professional ENGLISH as an organized study handout. Keep the scientific meaning, important details, headings, definitions, mechanisms, numbers, units, and medical/pharmaceutical terminology. Do not translate it to Arabic in this mode. Do not invent information.""",
    "translate_ar": """Translate the uploaded study material into accurate medical/pharmaceutical ARABIC. Keep the original English term beside the Arabic translation when useful. Preserve headings, definitions, mechanisms, numbers, units, and important details. Do not invent information.""",
    "study": """أنشئ خطة دراسة عملية لهذه المادة: ترتيب المواضيع، ماذا أحفظ، ماذا أفهم، ثم أسئلة ومراجعة. اجعلها مناسبة لطالب طب/صيدلة.""",
    "questions": """أنشئ أسئلة امتحانية من المادة، متنوعة بين MCQ وأسئلة فهم. اذكر الإجابة الصحيحة وسببها. لا تدّعِ أن أي سؤال مضمون في الامتحان.""",
    "flashcards": """أنشئ Flashcards تغطي التعاريف والمصطلحات والآليات والأرقام والمقارنات المهمة في المادة. كل بطاقة سؤال/مصطلح ثم جواب مختصر.""",
    "terms": """استخرج أهم المصطلحات الطبية والصيدلانية من المادة في جدول نصي: English | العربية | شرح مبسط.""",
    "revision": """أنشئ مراجعة سريعة جداً للمادة: أهم ما يجب حفظه، أهم ما يجب فهمه، الأخطاء الشائعة، ونقاط High-Yield.""",
    "progress": """قيّم بنية المادة وصعوبتها ووضوحها، ثم أعطِ خطة لتحسين الدراسة. لا تدّعي معرفة مستوى الطالب الحقيقي دون بيانات اختبار.""",
    "drug": """اعمل بطاقة تعليمية عن الدواء أو المادة الفعالة المذكورة في السؤال اعتماداً على المادة المرفوعة أولاً. اذكر: الاسم العلمي، الفئة الدوائية، آلية العمل إذا كانت موجودة، الاستعمالات العامة المذكورة، الأشكال الصيدلانية/التراكيز إن وردت، أهم التحذيرات أو الآثار الجانبية إن وردت. لا تعطِ جرعة شخصية أو وصفة علاجية. إذا لم تكن المعلومة في الملف فقل ذلك بوضوح، وأي معرفة إضافية ابدأها بعبارة "معلومة إضافية".""",
}


class FeatureRequest(BaseModel):
    feature: str
    query: str = ""

def _feature_instruction(feature: str, query: str = ""):
    instruction = FEATURE_PROMPTS.get(feature)
    if not instruction:
        raise HTTPException(400, "الميزة غير معروفة.")
    if query.strip():
        instruction += "\n\nموضوع/طلب الطالب: " + query.strip()
    return instruction

def _material_from_upload(data: bytes, filename: str, mime: str):
    text = extract_text(data, filename)
    image_data = base64.b64encode(data).decode("ascii") if mime.startswith("image/") else None
    if not text and not image_data:
        raise HTTPException(400, "لم أستطع قراءة هذا الملف. استخدم PDF/DOCX/PPTX/TXT أو صورة واضحة.")
    return text, image_data, mime

def ask_ai_with_optional_material(instruction: str, text: str = "", image_data: str | None = None, mime: str = "", max_output_tokens: int = 1600) -> str:
    require_client()
    if image_data:
        try:
            image_part = types.Part.from_bytes(data=base64.b64decode(image_data), mime_type=mime)
            contents = [instruction, image_part]
        except Exception as exc:
            raise RuntimeError("GEMINI_REQUEST_ERROR") from exc
    elif text:
        contents = instruction + "\n\nSTUDY MATERIAL PROVIDED FOR THIS REQUEST:\n" + text[:LARGE_FILE_CHARS]
    else:
        contents = instruction
    return _gemini_generate(contents, max_output_tokens=max_output_tokens)


@app.post("/feature")
def feature(body: FeatureRequest, request: Request):
    rate_limit(request)
    instruction=_feature_instruction(body.feature, body.query)
    try:
        return {"feature":body.feature,"answer":ask_ai_with_optional_material(instruction,max_output_tokens=1600)}
    except Exception as exc:
        msg=str(exc).lower()
        if "429" in msg or "rate_limit" in msg: raise HTTPException(429,"تم الوصول إلى حد Gemini مؤقتاً. انتظر ثم أعد المحاولة.")
        raise HTTPException(502,f"فشل تنفيذ الميزة: {exc}")

@app.post("/feature-file")
async def feature_file(request: Request, feature: str = "", query: str = "", file: UploadFile = File(...)):
    rate_limit(request)
    instruction=_feature_instruction(feature, query)
    data=await file.read()
    ensure_upload_size(data)
    text,image_data,mime=_material_from_upload(data,file.filename,file.content_type or "")
    try:
        answer=ask_ai_with_optional_material(instruction,text,image_data,mime,max_output_tokens=1800)
        return {"feature":feature,"filename":file.filename,"answer":answer}
    except Exception as exc:
        msg=str(exc).lower()
        if "429" in msg or "rate_limit" in msg: raise HTTPException(429,"تم الوصول إلى حد Gemini مؤقتاً. انتظر ثم أعد المحاولة.")
        raise HTTPException(502,f"فشل تنفيذ الميزة: {exc}")

class ChatRequest(BaseModel):
    question: str

@app.post("/chat")
def chat(body: ChatRequest, request: Request):
    rate_limit(request)
    q=body.question.strip()
    if not q: raise HTTPException(400,"اكتب السؤال.")
    instruction=f"""أجب عن سؤال الطالب بشكل مباشر وواضح. لا تفترض وجود ملف أو مادة دراسية إذا لم يتم إرفاقها في هذا الطلب. إذا كان السؤال طبياً أو صيدلانياً فاجعله تعليمياً ولا تعطِ وصفة شخصية أو جرعة شخصية.\n\nسؤال الطالب:\n{q}"""
    try:
        return {"answer":ask_ai_with_optional_material(instruction,max_output_tokens=800)}
    except Exception as exc:
        code=str(exc)
        if code == "GEMINI_AUTH_ERROR": raise HTTPException(502,"مفتاح Gemini غير صالح أو منتهي على Render.")
        if code == "GEMINI_RATE_LIMIT": raise HTTPException(429,"وصلنا إلى حد Gemini المجاني/المؤقت. انتظر قليلاً ثم جرّب مرة واحدة.")
        if code == "GEMINI_TIMEOUT": raise HTTPException(504,"Gemini لم يرد خلال المهلة. جرّب مرة واحدة بعد قليل.")
        raise HTTPException(502,"تعذر الوصول إلى Gemini حالياً. راجع سجل Render لمعرفة السبب.")

@app.post("/chat-file")
async def chat_file(request: Request, question: str = "", file: UploadFile = File(...)):
    rate_limit(request)
    q=question.strip()
    if not q: raise HTTPException(400,"اكتب السؤال.")
    data=await file.read()
    ensure_upload_size(data)
    text,image_data,mime=_material_from_upload(data,file.filename,file.content_type or "")
    instruction=f"""أجب عن سؤال الطالب اعتماداً على الملف المرفق أولاً. إذا لم تجد الإجابة في الملف، وضّح ذلك ثم استخدم معرفتك العامة عند الحاجة وسمها بوضوح: معلومة إضافية. لا تخترع معلومات. إذا كان السؤال طبياً أو صيدلانياً فاجعله تعليمياً ولا تعطِ وصفة شخصية أو جرعة شخصية.\n\nسؤال الطالب:\n{q}"""
    try:
        return {"answer":ask_ai_with_optional_material(instruction,text,image_data,mime,max_output_tokens=900),"filename":file.filename}
    except Exception as exc:
        code=str(exc)
        if code == "GEMINI_AUTH_ERROR": raise HTTPException(502,"مفتاح Gemini غير صالح أو منتهي على Render.")
        if code == "GEMINI_RATE_LIMIT": raise HTTPException(429,"وصلنا إلى حد Gemini المجاني/المؤقت. انتظر قليلاً ثم جرّب مرة واحدة.")
        if code == "GEMINI_TIMEOUT": raise HTTPException(504,"Gemini لم يرد خلال المهلة. جرّب مرة واحدة بعد قليل.")
        raise HTTPException(502,"تعذر الوصول إلى Gemini حالياً. راجع سجل Render لمعرفة السبب.")

@app.post("/reset")
def reset():
    last_material.update({"name": None, "text": "", "image_data": None, "mime": None})
    return {"ok": True}
